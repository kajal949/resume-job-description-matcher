import fitz
from docx import Document
from io import BytesIO


def extract_resume_text(uploaded_file):
    """Extract text from an uploaded PDF or DOCX resume."""

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith(".pdf"):
        text_parts = []

        with fitz.open(
            stream=file_bytes,
            filetype="pdf"
        ) as pdf:
            for page in pdf:
                text_parts.append(page.get_text())

        return "\n".join(text_parts).strip()

    elif file_name.endswith(".docx"):
        document = Document(BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        # Also extract text from tables in the resume.
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs).strip()

    else:
        raise ValueError(
            "Please upload a PDF or DOCX file."
        )